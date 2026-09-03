"""Regenerate tests/fixtures/eu_bulletin_slice.xlsx from the real bulletin.

The live workbook is ~4.5 MB, far too big to commit. This takes a slice of it
that keeps every layout feature the parser depends on:

- both price sheets, with their real header/unit rows
- column A dates, newest-first, plus the trailing blank + "Notes:" footer
- the FI, SE and IT country blocks: SE is the width-8 case (it carries an
  extra SE_exchange_rate column, so blocks are not a fixed width), IT has a
  genuinely missing week, FI is what the dashboard displays
- a per-tonne fuel-oil column, so the unit guard has something to reject

Usage (from the repo root, with the workbook downloaded already):

    python tests/fixtures/make_eu_fixture.py path/to/bulletin.xlsx

Committed output is what the tests read; rerun this only if the source layout
changes, and update docs/SCRAPER.md in the same pass.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

SHEETS = ("Prices with taxes", "Prices wo taxes")
COUNTRIES = ("FI", "SE", "IT")
DATA_ROWS = 6
OUT_PATH = Path(__file__).resolve().parent / "eu_bulletin_slice.xlsx"


def block_bounds(header_row, country_row):
    """(start, end) column indices of every CTR-marked country block."""
    starts = [i for i, v in enumerate(header_row) if v == "CTR"]
    return {
        country_row[s]: (s, e)
        for s, e in zip(starts, starts[1:] + [len(header_row)])
        if isinstance(country_row[s], str)
    }


def main(src: Path) -> None:
    source = openpyxl.load_workbook(src, read_only=True, data_only=True)
    out = openpyxl.Workbook()
    out.remove(out.active)

    for sheet_name in SHEETS:
        grid = list(source[sheet_name].iter_rows(values_only=True))
        header_row, unit_row, first_data = grid[0], grid[2], grid[3]
        bounds = block_bounds(header_row, first_data)

        keep = [0]  # column A: dates
        for country in COUNTRIES:
            start, end = bounds[f"{country}_"]
            keep.extend(range(start, end))

        target = out.create_sheet(sheet_name)
        for row in (grid[0], grid[1], grid[2], *grid[3 : 3 + DATA_ROWS]):
            target.append([row[i] if i < len(row) else None for i in keep])
        target.append([None] * len(keep))
        target.append(["Notes:"] + [None] * (len(keep) - 1))

    out.save(OUT_PATH)
    print(f"wrote {OUT_PATH} ({OUT_PATH.stat().st_size} bytes)")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(Path(sys.argv[1]))
