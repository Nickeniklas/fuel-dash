"""EU Weekly Oil Bulletin ingest: official weekly national prices into eu.db.

Second data source alongside the polttoaine.net scraper. Where the scraper
gives per-station Helsinki-area prices with no history before our first poll,
this gives one official national figure per week going back to 2005 — the
long-range context the crowdsourced data can't provide.

Writes to its own SQLite file, eu.db, rather than fuel.db -- see db.py's module
docstring for why.

Source: DG Energy's single "prices history" workbook, one stable URL holding
the whole series, refreshed weekly. There is no per-bulletin URL and no date
arithmetic on filenames: one file, one request. Layout facts are in
docs/SCRAPER.md ("EU Weekly Oil Bulletin"); code against that doc.

Licence: reproduction authorised provided the source is acknowledged. The
dashboard footer carries the attribution.
"""

from __future__ import annotations

import argparse
import io
import logging
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import requests

from db import connect_eu, ingest_eu_rows, latest_eu_week
from parser import USER_AGENT

logger = logging.getLogger(__name__)

BULLETIN_URL = (
    "https://energy.ec.europa.eu/document/download/"
    "906e60ca-8b6a-44e7-8589-652854d2fd3f_en"
    "?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx"
)

# Parsed and stored now even though the dashboard only displays FI — adding
# another country later is then a display change, not a re-ingest.
COUNTRIES = ("FI", "SE", "DE", "IT")

# Our fuel vocabulary -> the bulletin's column-name fragment. The bulletin has
# no 98E series at all (verified: euro95, diesel, heating oil, two fuel oils,
# LPG), so '98' has no entry here and the dashboard handles its absence.
FUEL_COLUMN_FRAGMENTS = {"95": "euro95", "dsl": "diesel"}

# Both tax variants live in this one workbook with an identical layout, so
# storing both costs one extra pass over rows we already have in memory.
SHEETS = {
    "Prices with taxes": ("price_with_tax", 1),
    "Prices wo taxes": ("price_wo_tax", 0),
}

# Row 1 holds the machine-readable column codes, row 2 human fuel names, row 3
# units. Data starts at row 4 and runs to a trailing "Notes:" footer, which is
# skipped by requiring a real date in column A rather than by row number.
HEADER_ROW = 1
UNIT_ROW = 3
EXPECTED_UNIT = "1000 l"

# The bulletin quotes EUR per 1000 litres; we store EUR/L like everything else.
LITRES_PER_UNIT = 1000

# Self-gate: skip the download entirely while our newest stored week is this
# fresh. The file is ~4.5 MB and the source only updates weekly, so without
# this the 12 h cron would refetch it ~14 times per new bulletin.
GATE_DAYS = 8

# Its own database, not fuel.db: the scraper's DB is committed on every 12 h
# poll and this series is both larger than it and only updated weekly.
EU_DB_PATH = Path(__file__).resolve().parent / "eu.db"


def download_workbook(url: str = BULLETIN_URL) -> bytes:
    """GET the bulletin workbook with the project's honest User-Agent."""
    logger.info("downloading %s", url)
    response = requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=120)
    response.raise_for_status()
    logger.info("downloaded %d bytes", len(response.content))
    return response.content


def _column_index(header_row: tuple, unit_row: tuple, code: str) -> int | None:
    """Locate a column by its row-1 code, verifying the row-3 unit.

    Country blocks are *not* a fixed width — non-euro countries carry an extra
    `<CC>_exchange_rate` column (their price columns are already converted to
    EUR) — so columns must be found by name, never by offset from the country
    marker. The unit check guards against a silent layout change moving us
    onto a per-tonne fuel-oil column.
    """
    for i, value in enumerate(header_row):
        if value == code:
            unit = unit_row[i] if i < len(unit_row) else None
            if unit != EXPECTED_UNIT:
                logger.warning(
                    "column %s has unit %r, expected %r — skipping",
                    code, unit, EXPECTED_UNIT,
                )
                return None
            return i
    logger.warning("column %s not found in the workbook header", code)
    return None


def parse_workbook(source: bytes | str | Path, countries: tuple[str, ...] = COUNTRIES) -> list[dict]:
    """Parse the bulletin workbook into a flat list of observation dicts.

    Each dict: {country, fuel, week_date (ISO), price (EUR/L), with_taxes}.
    Rows whose date cell isn't a real date (label rows, the trailing "Notes:"
    footer, blank padding) and cells with no numeric value are skipped.
    """
    stream = io.BytesIO(source) if isinstance(source, bytes) else source
    workbook = openpyxl.load_workbook(stream, read_only=True, data_only=True)

    rows: list[dict] = []
    for sheet_name, (fragment, with_taxes) in SHEETS.items():
        if sheet_name not in workbook.sheetnames:
            logger.warning("sheet %r missing from the workbook — skipping", sheet_name)
            continue
        sheet = workbook[sheet_name]
        grid = list(sheet.iter_rows(values_only=True))
        if len(grid) < UNIT_ROW:
            logger.warning("sheet %r has too few rows to hold a header", sheet_name)
            continue

        header_row, unit_row = grid[HEADER_ROW - 1], grid[UNIT_ROW - 1]
        columns = {}
        for country in countries:
            for fuel, fuel_fragment in FUEL_COLUMN_FRAGMENTS.items():
                code = f"{country}_{fragment}_{fuel_fragment}"
                index = _column_index(header_row, unit_row, code)
                if index is not None:
                    columns[(country, fuel)] = index
        if not columns:
            logger.warning("sheet %r: no requested columns found", sheet_name)
            continue

        for data_row in grid[UNIT_ROW:]:
            raw_date = data_row[0] if data_row else None
            if not isinstance(raw_date, (datetime, date)):
                continue  # label row, blank padding, or the "Notes:" footer
            resolved = raw_date.date() if isinstance(raw_date, datetime) else raw_date
            week_date = resolved.isoformat()
            for (country, fuel), index in columns.items():
                value = data_row[index] if index < len(data_row) else None
                if isinstance(value, bool) or not isinstance(value, (int, float)):
                    continue  # week not published for this country/fuel
                rows.append(
                    {
                        "country": country,
                        "fuel": fuel,
                        "week_date": week_date,
                        "price": value / LITRES_PER_UNIT,
                        "with_taxes": with_taxes,
                    }
                )
    workbook.close()
    return rows


def should_skip(conn, today: date | None = None, gate_days: int = GATE_DAYS) -> bool:
    """True if our newest stored week is fresh enough to skip the download."""
    if today is None:
        today = date.today()
    newest = latest_eu_week(conn)
    if newest is None:
        return False
    cutoff = today - timedelta(days=gate_days)
    if date.fromisoformat(newest) >= cutoff:
        logger.info(
            "newest stored week %s is within %d days of %s — skipping download",
            newest, gate_days, today.isoformat(),
        )
        return True
    return False


def update(
    eu_db_path: Path = EU_DB_PATH,
    url: str = BULLETIN_URL,
    today: date | None = None,
    source: bytes | str | Path | None = None,
) -> int:
    """Self-gated ingest: skip while our data is fresh, else fetch and upsert.

    Returns the number of rows stored (0 when gated out). Idempotent — the
    bulletin restates its whole history on every publish, so re-running just
    overwrites the same rows with the same values. Passing `source` parses a
    local file and bypasses both the download and the gate.
    """
    conn = connect_eu(eu_db_path)
    try:
        if source is None and should_skip(conn, today=today):
            return 0
        payload = source if source is not None else download_workbook(url)
        rows = parse_workbook(payload)
        stored = ingest_eu_rows(conn, rows)
        logger.info("parsed %d observations, stored %d", len(rows), stored)
        return stored
    finally:
        conn.close()


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    cli = argparse.ArgumentParser(description="Ingest the EU Weekly Oil Bulletin.")
    cli.add_argument(
        "--file",
        help="parse a local copy of the workbook instead of downloading "
             "(also bypasses the self-gate)",
    )
    args = cli.parse_args()
    logger.info("EU bulletin ingest with User-Agent: %s", USER_AGENT)
    update(source=args.file)
